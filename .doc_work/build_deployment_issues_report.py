from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "AWS_GCP_Firebase_Deployment_Issues.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_YELLOW = "FFF4CE"
PALE_RED = "FDE9E7"
PALE_GREEN = "E7F4E4"
MUTED = "666666"
BLACK = "000000"


def set_font(run, name="Calibri", size=11, bold=False, color=BLACK, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def format_cell_text(cell, bold=False, color=BLACK, size=9.5):
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            set_font(run, size=size, bold=bold, color=color)


def mark_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.5 + 0.25 * level)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    run = paragraph.add_run(text)
    set_font(run)
    return paragraph


def create_restarted_numbering(doc):
    numbering = doc.part.numbering_part.element
    style_num_id = doc.styles["List Number"]._element.pPr.numPr.numId.val
    base_num = numbering.find(f"w:num[@w:numId='{style_num_id}']", {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
    abstract_num_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    new_num_id = max(existing_ids, default=0) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_num_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return new_num_id


def add_number(doc, text, num_id=None):
    paragraph = doc.add_paragraph(style="List Number")
    if num_id is not None:
        num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_numId().val = num_id
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    run = paragraph.add_run(text)
    set_font(run)
    return paragraph


def add_callout(doc, label, text, fill=PALE_YELLOW):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.left_indent = Inches(0.15)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    label_run = paragraph.add_run(label + " ")
    set_font(label_run, bold=True, color=DARK_BLUE)
    body_run = paragraph.add_run(text)
    set_font(body_run)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
):
    style = styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(header.add_run("AiSteth Deployment Review"), size=9, bold=True, color=MUTED)
footer = section.footer.paragraphs[0]
add_page_number(footer)

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(18)
title.paragraph_format.space_after = Pt(4)
set_font(title.add_run("AWS, GCP, and Firebase Deployment Issues"), size=24, bold=True, color=BLACK)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
set_font(subtitle.add_run("Simple-English technical report"), size=14, color=MUTED)

for label, value in (
    ("System", "AiSteth dashboard frontend and FastAPI backend"),
    ("Cloud", "Google Cloud Run with AWS data services and Firebase Authentication"),
    ("Report date", "11 September 2026"),
    ("Purpose", "Explain the deployment problems, their causes, and the next actions"),
):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(label + ": "), bold=True)
    set_font(p.add_run(value))

add_callout(
    doc,
    "Main message:",
    "The containers can start, but login can still fail because authentication crosses three systems: Firebase checks the user token, AWS DynamoDB checks the user role, and Google Cloud Run carries the request between the frontend and backend.",
    PALE_YELLOW,
)

doc.add_heading("1. System in simple words", level=1)
intro = doc.add_paragraph()
intro.add_run("The frontend runs on Google Cloud Run. The backend also runs on Google Cloud Run. Firebase proves who the user is. AWS DynamoDB decides what that user can see. Elasticsearch supplies dashboard data. A login works only when every step works.")

add_number(doc, "The browser signs in with Firebase and receives a Firebase ID token.")
add_number(doc, "The frontend sends that token to the backend in the Authorization header.")
add_number(doc, "The backend asks Firebase Admin to verify the token.")
add_number(doc, "The backend reads the user's role from the AWS DynamoDB user-roles table.")
add_number(doc, "If the role is valid, the backend returns the user and dashboard data.")

doc.add_heading("2. Current issue summary", level=1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
headers = ["Area", "Problem", "Current status", "Priority"]
for index, text in enumerate(headers):
    table.rows[0].cells[index].text = text
    set_cell_shading(table.rows[0].cells[index], LIGHT_BLUE)
    format_cell_text(table.rows[0].cells[index], bold=True, color=DARK_BLUE)
mark_table_header(table.rows[0])

rows = [
    ("GCP", "Old Python package imports did not match the Docker folder layout.", "Fixed in code", "Done"),
    ("GCP", "Cloud Run Host header was rejected.", "Fixed in code/config", "Done"),
    ("GCP", "Frontend session call returned 503 because backend auth returned 503.", "Root cause found", "High"),
    ("AWS", "Deployed AWS key was rejected as an invalid security token.", "Confirmed", "Critical"),
    ("AWS", "Another tested IAM identity did not have DynamoDB Query permission.", "Confirmed", "Critical"),
    ("Firebase", "Backend now expects Firebase ID tokens, not old Cognito tokens.", "Must match frontend", "High"),
    ("Security", "Deployment environment file can enter the Docker build context.", "Needs protection", "Critical"),
]
for area, problem, state, priority in rows:
    cells = table.add_row().cells
    for index, text in enumerate((area, problem, state, priority)):
        cells[index].text = text
        format_cell_text(cells[index])
    if priority == "Critical":
        set_cell_shading(cells[3], PALE_RED)
    elif priority == "Done":
        set_cell_shading(cells[3], PALE_GREEN)
set_table_geometry(table, [1200, 4200, 2400, 1560])

doc.add_heading("3. Google Cloud Platform issues", level=1)
doc.add_heading("3.1 Python import failure during startup", level=2)
doc.add_paragraph(
    "The Docker image copies the backend directly into /app. There is no /app/backend_dashboard folder. Older imports used backend_dashboard.api and similar names. Cloud Run could not find that package, so the container stopped with ModuleNotFoundError. The imports were changed to match the real /app layout. The ASGI target is main:app."
)

doc.add_heading("3.2 Invalid Host header", level=2)
doc.add_paragraph(
    "TrustedHostMiddleware originally allowed only localhost and test hosts. A Cloud Run request uses a run.app hostname, so Starlette returned 'Invalid host header'. The allowlist was expanded for the production backend hostname and *.run.app. A Cloud Run environment variable can still override the code default, so the deployed value must stay correct."
)

doc.add_heading("3.3 Frontend session returns 503", level=2)
doc.add_paragraph(
    "Cloud logs showed a clear chain. The frontend received POST /api/auth/session. It then called backend GET /api/v1/auth/me. The backend returned 503, so the frontend returned 503 to the browser. The frontend was passing the failure through; it was not the original cause."
)

doc.add_heading("3.4 Environment deployment risk", level=2)
doc.add_paragraph(
    "Cloud Build deploys the backend with backend-env.yaml. This file contains deployment settings and may contain secrets. Docker also runs COPY . ., and the current .dockerignore does not clearly exclude backend-env.yaml. This can place secrets inside an image layer. The deployment file should be available to the deploy step but should not be copied into the final application image."
)
add_callout(doc, "Security action:", "Do not commit environment files containing passwords or cloud keys. Move secret values to Secret Manager and rotate any key that has been shared in chat, logs, source control, or an image.", PALE_RED)

doc.add_heading("4. AWS issues", level=1)
doc.add_heading("4.1 Invalid AWS security token", level=2)
doc.add_paragraph(
    "The backend uses boto3 to query the DynamoDB user-roles table. Cloud Run logs showed UnrecognizedClientException. AWS said the security token was invalid. This means the access-key pair used by the running revision was not accepted by AWS. Python cannot repair an invalid cloud key."
)

doc.add_heading("4.2 Missing DynamoDB Query permission", level=2)
doc.add_paragraph(
    "A second AWS identity was able to identify itself, but AWS denied dynamodb:Query on the dev_user_roles table. The backend catches this AWS ClientError and returns HTTP 503 with the safe message 'Authorization service is temporarily unavailable'."
)

doc.add_heading("4.3 Why local login looked successful", level=2)
doc.add_paragraph(
    "Local mode can bypass the DynamoDB role query when DASHBOARD_AUTHORIZATION_MODE is firebase-local and the request comes from a loopback address. Cloud Run requests are not loopback requests. The server therefore performs the real DynamoDB check and exposes AWS configuration or permission problems that local testing can hide."
)

doc.add_heading("4.4 AWS fix", level=2)
add_bullet(doc, "Use one active AWS credential source. Do not deploy duplicate access-key variables.")
add_bullet(doc, "Allow dynamodb:Query on the exact user-roles table used by the backend.")
add_bullet(doc, "Confirm that the table partition key is user_id, because the backend queries that key.")
add_bullet(doc, "Confirm that every allowed user has exactly one active role record.")
add_bullet(doc, "After the permission fix, test with a real Firebase user. No record should return 403, not 503.")

doc.add_heading("5. Firebase issues and checks", level=1)
doc.add_heading("5.1 Authentication system changed", level=2)
doc.add_paragraph(
    "The current backend code verifies Firebase ID tokens. Earlier deployment work used Cognito settings and Cognito tokens. The frontend and backend must use the same identity system. A Cognito token is not a Firebase token and will be rejected."
)

doc.add_heading("5.2 What the backend expects", level=2)
add_bullet(doc, "FIREBASE_PROJECT_ID must point to the Firebase project that issued the token.")
add_bullet(doc, "GOOGLE_CLOUD_PROJECT should point to the same Google Cloud project.")
add_bullet(doc, "The frontend must send the Firebase ID token as: Authorization: Bearer <token>.")
add_bullet(doc, "The email in the verified token must match user_id in the DynamoDB role table. If there is no email, the backend uses the Firebase UID.")
add_bullet(doc, "Cloud Run uses its Google service account through Application Default Credentials. A Firebase service-account JSON file should not be baked into the image.")

doc.add_heading("5.3 Expected Firebase failures", level=2)
firebase_table = doc.add_table(rows=1, cols=3)
firebase_table.style = "Table Grid"
for index, text in enumerate(("Case", "Backend result", "Meaning")):
    firebase_table.rows[0].cells[index].text = text
    set_cell_shading(firebase_table.rows[0].cells[index], LIGHT_GRAY)
    format_cell_text(firebase_table.rows[0].cells[index], bold=True, color=DARK_BLUE)
mark_table_header(firebase_table.rows[0])
for case, result, meaning in (
    ("No bearer token", "401", "The browser did not send a token."),
    ("Bad or expired Firebase token", "401", "Firebase could not validate the user token."),
    ("Firebase/Google service error", "503", "Token verification service was unavailable."),
    ("Valid token, no role record", "403", "User is known but has no dashboard permission."),
    ("Valid token, DynamoDB error", "503", "AWS role lookup failed."),
):
    cells = firebase_table.add_row().cells
    for index, text in enumerate((case, result, meaning)):
        cells[index].text = text
        format_cell_text(cells[index])
set_table_geometry(firebase_table, [3000, 1200, 5160])

doc.add_heading("5.4 Swagger limitation", level=2)
doc.add_paragraph(
    "The current backend Swagger scheme accepts a Firebase ID token. Firebase email/password login normally happens in the frontend with the Firebase client SDK. Swagger does not automatically become a Firebase login page. For backend API testing, obtain a Firebase ID token from a trusted test client and paste it into Swagger Authorize."
)

doc.add_heading("6. Recommended action plan", level=1)
action_num_id = create_restarted_numbering(doc)
add_number(doc, "Rotate every AWS key that has been exposed and remove old or invalid keys from Cloud Run.", action_num_id)
add_number(doc, "Give the production AWS identity only the DynamoDB Query permission needed for the user-roles table.", action_num_id)
add_number(doc, "Store AWS and Elasticsearch secrets in Secret Manager, not in source files or Docker image layers.", action_num_id)
add_number(doc, "Confirm that frontend login uses Firebase and that it sends a Firebase ID token to the backend.", action_num_id)
add_number(doc, "Confirm that Firebase project IDs match in frontend, backend, and Cloud Run.", action_num_id)
add_number(doc, "Add the production frontend URL to DASHBOARD_CORS_ORIGINS and keep the backend run.app hosts in DASHBOARD_ALLOWED_HOSTS.", action_num_id)
add_number(doc, "Deploy a new backend revision and test /health, /docs, /openapi.json, and /api/v1/auth/me.", action_num_id)
add_number(doc, "Check Cloud Run stderr logs. A successful fix must remove UnrecognizedClientException and AccessDeniedException from the auth path.", action_num_id)

doc.add_heading("7. Go-live checklist", level=1)
checks = [
    "Backend container starts on 0.0.0.0 and uses PORT=8080.",
    "Health, docs, and OpenAPI return HTTP 200.",
    "Production Host header is accepted.",
    "Browser preflight request receives the correct CORS headers.",
    "Firebase login returns an ID token for the correct project.",
    "Backend accepts that token and returns 401 for an invalid token.",
    "DynamoDB Query succeeds with the production AWS identity.",
    "A valid user has exactly one active role record.",
    "Frontend session returns 200 and does not hide a backend 503.",
    "No secret values are present in Git history, logs, image layers, or the final report.",
]
for item in checks:
    add_bullet(doc, "Check: " + item)

doc.add_heading("8. Final conclusion", level=1)
doc.add_paragraph(
    "The main deployment problem is not one cloud by itself. The application crosses GCP, Firebase, and AWS. GCP can run the container correctly while login still fails in AWS. Firebase can verify the user correctly while DynamoDB still rejects the role lookup. The safest fix is to keep Firebase for identity, use one valid and limited AWS identity for role reads, and keep all secrets outside the Docker image."
)

add_callout(doc, "Definition of done:", "A real production user signs in with Firebase, the backend verifies the Firebase ID token, DynamoDB returns one active role, /api/v1/auth/me returns 200, and the frontend session also returns 200.", PALE_GREEN)

doc.core_properties.title = "AWS, GCP, and Firebase Deployment Issues"
doc.core_properties.subject = "Simple-English deployment issue report"
doc.core_properties.author = "AiSteth Engineering"
doc.core_properties.keywords = "AWS, GCP, Cloud Run, Firebase, DynamoDB, deployment"
doc.save(OUTPUT)
print(OUTPUT)
